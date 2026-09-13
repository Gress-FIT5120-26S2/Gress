from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_PATH = Path(__file__).with_name("KitchMemo_Test_Cases_and_Report.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="173D31")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color="173D31", size=18)
SECTION_FONT = Font(name="Calibri", bold=True, color="173D31", size=13)
BODY_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Border(
    left=Side(style="thin", color="D9E2DE"),
    right=Side(style="thin", color="D9E2DE"),
    top=Side(style="thin", color="D9E2DE"),
    bottom=Side(style="thin", color="D9E2DE"),
)
ALT_FILL = PatternFill("solid", fgColor="F7F9F7")
HIGH_FILL = PatternFill("solid", fgColor="FDE8D8")
MED_FILL = PatternFill("solid", fgColor="FFF6E5")
LOW_FILL = PatternFill("solid", fgColor="EAF6F1")
PASS_FILL = PatternFill("solid", fgColor="D9F2E4")
FAIL_FILL = PatternFill("solid", fgColor="F8D7DA")
BLOCK_FILL = PatternFill("solid", fgColor="FCE8C8")
NA_FILL = PatternFill("solid", fgColor="E9EEF1")
LABEL_FILL = PatternFill("solid", fgColor="E8F0ED")

HEADERS = [
    "Case ID",
    "Module",
    "Title",
    "Priority",
    "Type",
    "Preconditions",
    "Test Steps",
    "Expected Result",
    "Actual Result",
    "Status",
    "Notes",
]

# Course-style cases covering completed KitchMemo features only.
CASES = [
    # Onboarding
    ("TC-ONB-001", "Onboarding", "Opening animation completes and reveals the kitchen", "High", "UI",
     "Fresh install or app relaunch; network available.",
     "1. Launch the app.\n2. Wait for the opening sequence.\n3. Observe the home kitchen.",
     "The opener finishes, the 3D kitchen appears, and the app is interactive. It does not stay on the opener forever.",
     "", "", ""),
    ("TC-ONB-002", "Onboarding", "First-use journey is shown only once", "High", "Functional",
     "Device has no completed first-use flag.",
     "1. Launch after a fresh install.\n2. Complete the first-use pages and enter the kitchen.\n3. Force-close and reopen the app.",
     "The journey appears on first launch. After completion it does not appear again on the same device.",
     "", "", ""),
    ("TC-ONB-003", "Onboarding", "First-use pages can be navigated forward and back", "Medium", "UI",
     "First-use journey is visible.",
     "1. Open the journey.\n2. Use Next through each page.\n3. Use Back to return to an earlier page.",
     "Progress text updates. Back restores the previous page. Enter kitchen completes the flow.",
     "", "", ""),
    ("TC-ONB-004", "Onboarding", "Device bootstrap creates a personal fridge", "High", "Integration",
     "New device ID; Express and Supabase reachable.",
     "1. Complete first launch.\n2. Open Fridge.\n3. Confirm the fridge name/scope.",
     "The device receives a personal fridge with default categories. Inventory can load without a no-fridge error.",
     "", "", ""),

    # Home
    ("TC-HOME-001", "Home", "Bottom tabs expose Home, Cart, Fridge, and Me", "High", "UI",
     "First-use journey completed.",
     "1. Look at the floating tab bar.\n2. Tap each visible tab.",
     "Home, Cart, Fridge, and Me (profile) are shown. Achievements / Wins is still hidden until its metrics API exists.",
     "", "", ""),
    ("TC-HOME-002", "Home", "3D fridge hotspot opens the unfiltered fridge", "High", "Functional",
     "Home kitchen is loaded; inventory may contain expiring items.",
     "1. Tap the white hotspot on the 3D fridge.\n2. Check the fridge filter chips.",
     "The fridge page opens. No filter chip is pre-selected. All active items are listed.",
     "", "", ""),
    ("TC-HOME-003", "Home", "3D cart hotspot opens shopping", "High", "Functional",
     "Home kitchen is loaded.",
     "1. Tap the shopping cart in the 3D scene.",
     "The shopping screen opens after the cinematic transition.",
     "", "", ""),
    ("TC-HOME-004", "Home", "3D mailbox opens the notification inbox", "High", "Functional",
     "Home kitchen is loaded.",
     "1. Tap the mailbox in the 3D scene.",
     "The notification inbox opens. Unread count matches the home mailbox badge.",
     "", "", ""),
    ("TC-HOME-005", "Home", "Stove and recipe book stay in-scene", "Medium", "UI",
     "Home kitchen is loaded.",
     "1. Tap the stove.\n2. Tap the recipe book.",
     "Camera/scene feedback occurs in place. The app does not navigate to a recipe destination.",
     "", "", ""),
    ("TC-HOME-006", "Home", "Expiring headline count matches fridge Expiring chip", "High", "Functional",
     "Fridge contains N active items that expire within 3 days and are not yet expired.",
     "1. Note the home headline count.\n2. Open Fridge from the tab bar.\n3. Compare the Expiring chip count.",
     "Home count equals the fridge Expiring chip count. It is not a hardcoded sample such as 2.",
     "", "", ""),
    ("TC-HOME-007", "Home", "Tapping the expiring headline opens fridge Expiring filter", "High", "Functional",
     "At least one expiring item exists.",
     "1. Tap the home headline (e.g. '3 ingredients to use soon').\n2. Inspect fridge chips and list.",
     "Fridge opens with the Expiring chip selected. The list shows only not-expired items with days left <= 3.",
     "", "", ""),
    ("TC-HOME-008", "Home", "Use first appears when nothing is expiring", "High", "Functional",
     "Fridge has no expiring items (empty fridge, or all items expire after 3 days / already expired).",
     "1. Open Home.\n2. Read the freshness headline.\n3. Tap Use first.",
     "Headline shows Use first (ZH: 先用) in a larger type size. Tap still opens fridge with the Expiring chip selected (empty list is allowed).",
     "", "", ""),
    ("TC-HOME-009", "Home", "Expired items are not counted as expiring", "High", "Functional",
     "One item expired yesterday; one item expires in 2 days.",
     "1. Check home count.\n2. Check fridge Expiring and Expired chips.",
     "Home count is 1. Expiring chip is 1. Expired chip includes the yesterday item. Expired items are not in Expiring.",
     "", "", ""),
    ("TC-HOME-010", "Home", "Items without expiry are not counted as expiring", "Medium", "Functional",
     "An active item has no expiry date.",
     "1. Open Home and Fridge Expiring.",
     "The undated item is absent from the expiring count and Expiring list.",
     "", "", ""),
    ("TC-HOME-011", "Home", "Mailbox numeric badge uses badgeCount, not raw unreadCount", "High", "Functional",
     "Inbox has unread items; notification badges are enabled and quiet hours are off.",
     "1. Note the orange number on the home mail button.\n2. Open the inbox and compare unread summary.\n3. Disable home badges in Me > Notifications & reminders.",
     "When badges are on, the orange number equals API badgeCount (capped at 99+). Turning badges off or entering quiet hours hides the number even if unread items remain. The mail icon can still show an unread glyph from unreadCount.",
     "", "", ""),
    ("TC-HOME-012", "Home", "Kitchen lighting follows time of day", "Low", "UI",
     "Device clock is at night, then daytime if possible.",
     "1. Observe home lighting and period label.",
     "Period copy matches the current phase (night/dawn/day/sunset). Night uses light status bar treatment.",
     "", "", ""),
    ("TC-HOME-013", "Home", "Tapping the home mailbox opens the full notification page", "High", "Functional",
     "Home kitchen is loaded.",
     "1. Tap the overlay mail button.",
     "A full notification centre page opens (back, title, refresh). It is not an inline placeholder under generic screen copy.",
     "", "", ""),

    # Fridge
    ("TC-FRG-001", "Fridge", "Inventory list loads active batches", "High", "Functional",
     "Fridge contains at least two active batches.",
     "1. Open Fridge from the tab bar.\n2. Pull to refresh.",
     "Active batches appear with name, quantity, storage, and freshness. Soft-deleted items are not listed.",
     "", "", ""),
    ("TC-FRG-002", "Fridge", "Tab-bar fridge entry is unfiltered", "High", "Functional",
     "Previous visit used the home expiring link.",
     "1. From Home, tap Use first / expiring headline.\n2. Go Home.\n3. Tap the Fridge tab.",
     "The second visit is unfiltered. The Expiring chip is not still selected.",
     "", "", ""),
    ("TC-FRG-003", "Fridge", "Search filters by food name", "Medium", "Functional",
     "Fridge contains Milk and Tomato.",
     "1. Type 'tom' in search.\n2. Clear search.",
     "Only matching names remain. Clearing restores the previous filter/category combination.",
     "", "", ""),
    ("TC-FRG-004", "Fridge", "Storage chips filter chilled, frozen, and pantry", "High", "Functional",
     "At least one item in each storage zone, if possible.",
     "1. Tap Chilled, Frozen, then Pantry.\n2. Confirm counts on chips.",
     "Each chip shows only items in that storage zone. Chip counts match the visible matching items.",
     "", "", ""),
    ("TC-FRG-005", "Fridge", "Expiring chip uses the 3-day not-expired rule", "High", "Functional",
     "Items at: expired, 0-3 days left, 4+ days left, no date.",
     "1. Select Expiring.",
     "Only not-expired items with daysLeft <= 3 appear. Expired, undated, and 4+ day items are excluded.",
     "", "", ""),
    ("TC-FRG-006", "Fridge", "Expired chip shows past-expiry active items", "High", "Functional",
     "At least one active item with expiresAt in the past.",
     "1. Select Expired.",
     "Only items whose expiry timestamp is before now appear.",
     "", "", ""),
    ("TC-FRG-007", "Fridge", "Restock chip shows items at or below the restock minimum", "High", "Functional",
     "A restock rule exists and remaining quantity is at or below minimum.",
     "1. Select Restock.",
     "Only items flagged needsRestock appear. Chip count matches.",
     "", "", ""),
    ("TC-FRG-008", "Fridge", "Category chips can combine with a status filter", "Medium", "Functional",
     "Expiring items exist in more than one category.",
     "1. Select Expiring.\n2. Select Vegetables.",
     "The list is the intersection of both conditions. Clearing filters restores all items.",
     "", "", ""),
    ("TC-FRG-009", "Fridge", "Empty expiring state is understandable", "Medium", "UI",
     "No expiring items.",
     "1. Open Fridge via the home Use first link.",
     "Expiring is selected and an empty state is shown. The app does not crash.",
     "", "", ""),
    ("TC-FRG-010", "Fridge", "Fridge assistant can be opened and closed", "Medium", "Functional",
     "Fridge has dated inventory.",
     "1. Open the fridge assistant.\n2. Ask a use-first / restock style question.\n3. Close it.",
     "Assistant answers from current batches. Closing returns to the fridge list.",
     "", "", ""),

    # Add method sheet
    ("TC-ADD-001", "Add item", "Add sheet offers manual entry and photo recognition", "High", "UI",
     "Fridge screen is open.",
     "1. Tap the add (+) control.",
     "A bottom sheet titled Add ingredients appears with Manual entry and Photo recognition cards. Photo recognition shows Recommended on its own top row.",
     "", "", ""),
    ("TC-ADD-002", "Add item", "Manual entry and Photo recognition titles align", "Medium", "UI",
     "Add sheet is open.",
     "1. Compare the two cards.",
     "Recommended sits above the camera icon, not over the title. Manual entry has matching top spacing so both titles start at the same height. No arrow buttons.",
     "", "", ""),
    ("TC-ADD-003", "Add item", "Closing the sheet without a choice returns to fridge", "Medium", "Functional",
     "Add sheet is open.",
     "1. Tap the close (X) or the backdrop.",
     "The sheet dismisses. Fridge remains visible. No camera or form opens.",
     "", "", ""),
    ("TC-ADD-004", "Add item", "Selecting manual entry opens the shared form", "High", "Functional",
     "Add sheet is open.",
     "1. Tap Manual entry.",
     "InventoryEntryFlow opens. The previous sheet does not intercept touches.",
     "", "", ""),
    ("TC-ADD-005", "Add item", "Selecting photo recognition opens the camera flow", "High", "Functional",
     "Camera permission can be granted.",
     "1. Tap Photo recognition.\n2. Grant permission if asked.",
     "The recognition camera / gallery flow opens.",
     "", "", ""),

    # Manual inventory
    ("TC-INV-001", "Inventory entry", "Manual save creates a new independent batch", "High", "Functional",
     "Form is open; a food name, quantity, storage, and optional expiry are filled.",
     "1. Enter Milk, quantity 2, chilled, expiry in 5 days.\n2. Save.\n3. Repeat with the same name on another day.",
     "Each save creates a separate batch. Same-name items are not merged.",
     "", "", ""),
    ("TC-INV-002", "Inventory entry", "Known food name loads a preset suggestion", "High", "Functional",
     "Seed presets exist (e.g. tomato, milk).",
     "1. Type tomato or a known alias.\n2. Wait for suggestion.\n3. Accept and save.",
     "Suggested storage, category, shelf life, and icon appear. Saving stores presetUid when the preset is enabled.",
     "", "", ""),
    ("TC-INV-003", "Inventory entry", "Unknown name offers AI generate instead of web search", "High", "Functional",
     "Gemini/Cloudflare credentials configured; name is not in presets.",
     "1. Type an uncommon food name.\n2. Tap AI generate.\n3. Review editable suggestion and save.",
     "A reusable global preset and icon are generated. The form stays editable. Inventory is written only after Save.",
     "", "", ""),
    ("TC-INV-004", "Inventory entry", "AI generate is rate-limited per device", "Medium", "Negative",
     "Server generation limit is 5 new foods per device per hour.",
     "1. Trigger AI generate more than 5 times in one hour.",
     "Further generate requests are rejected with a clear error. Manual entry still works.",
     "", "", ""),
    ("TC-INV-005", "Inventory entry", "Expiry cannot be earlier than stocked time", "Medium", "Negative",
     "Manual form is open.",
     "1. Set expiry before the recorded stocked date/time.\n2. Attempt save.",
     "Save is rejected or blocked. No invalid batch is created.",
     "", "", ""),
    ("TC-INV-006", "Inventory entry", "Quantity must be greater than zero", "Medium", "Negative",
     "Manual form is open.",
     "1. Set quantity to 0 or a negative number.\n2. Attempt save.",
     "Save is rejected. Remaining quantity stays within 0 to initial quantity after create.",
     "", "", ""),
    ("TC-INV-007", "Inventory entry", "Optional restock rule is saved with the batch", "Medium", "Functional",
     "Form supports restock thresholds.",
     "1. Enable restock, set minimum and a higher target.\n2. Save.\n3. Open the item and Restock chip.",
     "Rule persists. When remaining quantity is at or below minimum, the item appears under Restock.",
     "", "", ""),
    ("TC-INV-008", "Inventory entry", "Form error from the API is visible", "Medium", "Negative",
     "Express is running; force a validation failure if possible.",
     "1. Submit invalid data.\n2. Read the form error.",
     "A real error message is shown. The form does not fail silently.",
     "", "", ""),

    # Photo recognition
    ("TC-REC-001", "Photo recognition", "Supported produce is recognised and prefilled", "High", "Functional",
     "Camera/gallery available; photo of tomato, banana, or another supported food.",
     "1. Capture or pick a photo.\n2. Confirm the review screen.\n3. Continue to the entry form without saving yet.",
     "Name is prefilled from recognition. Freshness maps shelf life: fresh = full preset days, semi_fresh = ceil(40%), rotten = shortest check window. Nothing is written until Save.",
     "", "", ""),
    ("TC-REC-002", "Photo recognition", "Unsupported or failed recognition falls back to manual entry", "High", "Negative",
     "Photo of an unsupported object, or API timeout.",
     "1. Submit the photo.\n2. Observe the result.",
     "User can retake or continue to a blank/manual form. The app does not crash. Image is not stored in Supabase.",
     "", "", ""),
    ("TC-REC-003", "Photo recognition", "User can edit prefilled values before save", "High", "Functional",
     "Recognition returned a draft.",
     "1. Change name, quantity, or expiry.\n2. Save.",
     "Saved batch uses the edited values, not the raw model output.",
     "", "", ""),
    ("TC-REC-004", "Photo recognition", "Oversized image is rejected", "Low", "Negative",
     "Image larger than 10 MB if testable.",
     "1. Select an oversized file.",
     "Request is rejected with a size/type error. No disk write of the photo occurs on the server.",
     "", "", ""),

    # Item detail
    ("TC-DET-001", "Inventory detail", "Quantity can be decreased and increased", "High", "Functional",
     "An active batch with remaining quantity 3 and initial 3.",
     "1. Open the item sheet.\n2. Decrease to 1.\n3. Increase back toward initial.",
     "Quantity updates immediately. Value cannot exceed initial quantity or go below 0. An inventory event is recorded.",
     "", "", ""),
    ("TC-DET-002", "Inventory detail", "Quantity 0 marks the batch consumed", "High", "Functional",
     "Active batch with remaining > 0.",
     "1. Set remaining quantity to 0.\n2. Return to the list.",
     "Batch lifecycle becomes consumed and it leaves the active list. Increasing from 0 can restore active.",
     "", "", ""),
    ("TC-DET-003", "Inventory detail", "Edit details updates name, storage, and expiry", "High", "Functional",
     "Active batch exists.",
     "1. Open edit.\n2. Change storage zone and expiry.\n3. Save.",
     "List card reflects the new storage and freshness. Version conflict returns 409 rather than silent overwrite.",
     "", "", ""),
    ("TC-DET-004", "Inventory detail", "Soft delete archives the batch", "High", "Functional",
     "Active batch exists.",
     "1. Delete / archive from the detail sheet.\n2. Refresh the fridge list.",
     "Item disappears from active inventory. Remaining quantity is zeroed. History is kept; it is not a hard delete.",
     "", "", ""),
    ("TC-DET-005", "Inventory detail", "Restock rule can be updated from the sheet", "Medium", "Functional",
     "Item is open.",
     "1. Set minimum/target.\n2. Save.\n3. Lower quantity to the minimum.",
     "needsRestock becomes true and the Restock chip includes the item.",
     "", "", ""),

    # Shopping
    ("TC-SHOP-001", "Shopping", "Restock tab lists suggested buys", "High", "Functional",
     "At least one enabled restock rule with stock at or below minimum.",
     "1. Open Cart tab.\n2. Open the Restock view.",
     "Suggested items match restock rules and current stock. Pull-to-refresh reloads.",
     "", "", ""),
    ("TC-SHOP-002", "Shopping", "Cart tab can add a manual item", "High", "Functional",
     "Shopping screen is open.",
     "1. Add an item via the shopping add sheet.\n2. Confirm it appears in Cart.",
     "Cart item is created against the current fridge. Quantity is editable.",
     "", "", ""),
    ("TC-SHOP-003", "Shopping", "Cart quantity can be changed and deleted", "High", "Functional",
     "Cart has one item.",
     "1. Use − / +.\n2. Type a quantity.\n3. Delete the row.",
     "Quantity updates. Delete removes the item. Shared members see the same cart after sync.",
     "", "", ""),
    ("TC-SHOP-004", "Shopping", "Adding a name that already exists in the fridge warns of a duplicate", "High", "Functional",
     "Fridge already contains Milk; cart add flow is open.",
     "1. Add Milk to the cart.\n2. Read the duplicate warning.",
     "A possible-duplicate warning is shown using current inventory names. User can still add if they choose.",
     "", "", ""),
    ("TC-SHOP-005", "Shopping", "Inventory peek shows what is already at home", "Medium", "Functional",
     "Fridge has inventory; shopping screen is open.",
     "1. Open inventory peek.\n2. Close it.",
     "Current fridge stock is visible without leaving shopping. Closing returns to cart/restock.",
     "", "", ""),
    ("TC-SHOP-006", "Shopping", "Checkout only includes checked cart items", "High", "Functional",
     "Cart has two items; only one is checked.",
     "1. Confirm checkout button count.\n2. Open checkout review.",
     "Checkout is offered only when at least one item is checked. Review lists only checked items. Unchecked items stay in the cart and are not stocked.",
     "", "", ""),
    ("TC-SHOP-007", "Shopping", "Shopping add sheet reuses the same method picker", "Medium", "UI",
     "Shopping add is opened.",
     "1. Inspect manual vs photo options.",
     "The shared AddItemMethodSheet is used. Layout matches fridge add (no arrows; Recommended on its own top row).",
     "", "", ""),
    ("TC-SHOP-008", "Shopping", "Checked checkout items can be stocked into the fridge", "High", "Functional",
     "One cart item is checked.",
     "1. Open checkout.\n2. Complete InventoryEntryFlow for that item.\n3. Open Fridge.",
     "The stocked item becomes an inventory batch and is removed from the cart. Fridge list updates.",
     "", "", ""),
    ("TC-SHOP-009", "Shopping", "Leaving checkout with unstocked checked items prompts first", "High", "Functional",
     "Checkout is open with at least one checked item not yet stocked.",
     "1. Attempt to close checkout.\n2. Choose stay, then leave anyway.",
     "A leave prompt appears. Stay keeps the review open. Leave anyway closes it without stocking remaining items.",
     "", "", ""),

    # Notifications
    ("TC-NOT-001", "Notifications", "Opening the inbox syncs expiring, expired, and restock events", "High", "Functional",
     "Fridge has an expiring item, an expired item, and a restock-needed item.",
     "1. Open the mailbox/inbox.",
     "Notifications are generated from live stock. Types include expiring, expired, and restock. No duplicate rows for the same batch/type (dedupe_key).",
     "", "", ""),
    ("TC-NOT-002", "Notifications", "Expiring notification uses the same 3-day window", "High", "Functional",
     "Item expires in 2 days; another expires in 5 days.",
     "1. Open inbox.\n2. Compare with fridge Expiring.",
     "Only the 2-day item creates an expiring notification. The 5-day item does not.",
     "", "", ""),
    ("TC-NOT-003", "Notifications", "Mark as read is per device", "High", "Functional",
     "Two devices in the same shared fridge, or simulate two device IDs.",
     "1. On device A, mark a notification read.\n2. Open inbox on device B.",
     "Device A shows it read. Device B still shows unread. Home badge on each device is independent.",
     "", "", ""),
    ("TC-NOT-004", "Notifications", "Opening a notification marks it read and updates counts", "High", "Functional",
     "At least one unread notification; badges enabled.",
     "1. Open the inbox.\n2. Tap an unread row.\n3. Close the detail sheet and return Home.",
     "A detail sheet shows localised title, body, and time. The row becomes read. unreadCount and badgeCount decrease. Home orange badge updates.",
     "", "", ""),
    ("TC-NOT-005", "Notifications", "Archived batches do not keep active inventory notifications visible", "Medium", "Functional",
     "A notification exists for a batch that is then archived.",
     "1. Archive the related item.\n2. Reopen inbox.",
     "Notifications tied to inactive batches are filtered out of the visible list.",
     "", "", ""),
    ("TC-NOT-006", "Notifications", "Inbox refresh and empty/error states work", "Medium", "UI",
     "Notification page is open.",
     "1. Pull/tap refresh.\n2. If possible, stop the API and refresh.\n3. Restore API on an empty fridge.",
     "Refresh reloads the snapshot. Failure shows a retry empty state. A fridge with no events shows the empty copy, not a crash.",
     "", "", ""),
    ("TC-NOT-007", "Notifications", "Shared inventory events appear for other members only", "High", "Functional",
     "Devices A and B share a fridge. Both have shared notifications enabled.",
     "1. On A, add, edit, or remove an item.\n2. Open B's inbox.\n3. Open A's inbox.",
     "B receives a shared notification with A's nickname and the food name. A does not receive a shared event for their own action. Expiring/expired/restock sync still runs from live stock.",
     "", "", ""),
    ("TC-NOT-008", "Notifications", "Category switches hide types from the inbox list", "High", "Functional",
     "Inbox contains expiring and restock items.",
     "1. In Me > Notifications, turn off freshness (expiring/expired).\n2. Reopen the inbox.",
     "Expiring/expired rows disappear. Restock remains if that switch is on. Turning the master switch off hides reminder history types according to server filtering. History is not marked read.",
     "", "", ""),
    ("TC-NOT-009", "Notifications", "Tapping a system notification opens the matching inbox item", "Medium", "Functional",
     "A scheduled/local or push notification with screen=notifications is delivered.",
     "1. Tap the OS notification card.",
     "The app opens the notification centre and, when an id is present, auto-opens that item.",
     "", "", ""),

    # Sharing
    ("TC-SHR-001", "Sharing", "Personal fridge can be named and switched to shared", "High", "Functional",
     "Device is on a personal fridge.",
     "1. Open the fridge space menu.\n2. Create / enable a household fridge with a custom name.",
     "Current fridge becomes shared (not a second container). An active invite code is generated.",
     "", "", ""),
    ("TC-SHR-002", "Sharing", "Invite code can be copied, shared, and shown as QR", "Medium", "Functional",
     "Fridge is shared.",
     "1. Open the share page.\n2. Copy code, system-share, and display QR.",
     "Code and QR represent the current active invite. Regenerating revokes the old code.",
     "", "", ""),
    ("TC-SHR-003", "Sharing", "Second personal device can join with a valid code", "High", "Functional",
     "Device B is a single-member personal fridge with some inventory.",
     "1. On B, join using A's active code.\n2. Open Fridge on both devices.",
     "Context returns the current invite plus member nicknames, avatar tokens, and order. Real device_id values are not shown. B's personal data merges into A's fridge. Same-name batches stay independent. Both devices see merged inventory.",
     "", "", ""),
    ("TC-SHR-004", "Sharing", "Join fails with distinct errors for expired, used, revoked, and missing codes", "High", "Negative",
     "Prepare expired, used, revoked, and random codes.",
     "1. Attempt join with each code.",
     "UI shows the matching message: expired, used, revoked, or not found/invalid. It does not collapse all failures into one generic error.",
     "", "", ""),
    ("TC-SHR-005", "Sharing", "Owner leaving takes only their owned active batches", "High", "Functional",
     "Shared fridge with items owned by A and B.",
     "1. Device A leaves.\n2. Check A's new personal fridge and B's remaining household fridge.",
     "A receives a new personal fridge with batches it owns. B keeps other members' items, achievements, and historical batches.",
     "", "", ""),
    ("TC-SHR-006", "Sharing", "Fridge can be renamed in shared mode", "Low", "Functional",
     "Shared fridge is active.",
     "1. Rename the fridge.\n2. Recheck the header on both devices.",
     "The new name is visible to members after refresh/sync.",
     "", "", ""),
    ("TC-SHR-007", "Sharing", "Shared notifications use the actor nickname, not device id", "High", "Functional",
     "Device A has display name 'Alex'. Device B is a member.",
     "1. A adds an item.\n2. Read B's shared notification copy.",
     "The title uses Alex (or the current language fallback actor) and the food name. It does not expose the raw device_id.",
     "", "", ""),

    # Profile
    ("TC-PRF-001", "Profile", "Me tab loads device nickname and fridge summary", "High", "Functional",
     "Device is bootstrapped.",
     "1. Open Me.\n2. Read greeting and fridge line.",
     "Greeting uses the saved display name (default if unset). Fridge line shows personal or shared name and member count. Avatar is a product token/colour, not a user photo.",
     "", "", ""),
    ("TC-PRF-002", "Profile", "Nickname can be edited between 1 and 32 characters", "High", "Functional",
     "Profile edit modal is available.",
     "1. Open edit profile.\n2. Save a valid 1–32 character name.\n3. Attempt empty or 33+ characters.",
     "Valid names persist via PATCH /api/profile and appear in greeting and shared notifications. Invalid lengths are rejected.",
     "", "", ""),
    ("TC-PRF-003", "Profile", "Taste preferences remain a non-persisting placeholder", "Low", "UI",
     "Me tab is open.",
     "1. Tap the taste/diet row.",
     "An informational alert is shown. No taste data is written to the server.",
     "", "", ""),
    ("TC-PRF-004", "Profile", "Language settings are reachable from Me", "High", "Functional",
     "Me tab is open.",
     "1. Open Language.\n2. Switch English/Chinese.\n3. Relaunch.",
     "The whole UI switches, including Home Use first / 先用. Preference persists.",
     "", "", ""),
    ("TC-PRF-005", "Profile", "Notifications & reminders and inbox history have separate entries", "High", "Functional",
     "Me tab is open.",
     "1. Open Notifications & reminders.\n2. Use View notification history.\n3. Return and confirm settings are still the settings page.",
     "Settings is a full-screen device preference page. History opens the notification centre. They are not the same screen.",
     "", "", ""),
    ("TC-PRF-006", "Profile", "Recovery and privacy rows open their sheets", "Medium", "Functional",
     "Me tab is open.",
     "1. Open device recovery.\n2. Open privacy.",
     "Recovery shows the configured/not-configured state. Privacy content opens. Raw Device-Credential is not displayed in logs or API responses.",
     "", "", ""),

    # Notification preferences
    ("TC-PREF-001", "Notification settings", "Master switch saves immediately for this device", "High", "Functional",
     "Me > Notifications & reminders is open.",
     "1. Toggle the master reminders switch off, then on.",
     "PATCH /api/notification-preferences succeeds. Delivery and type switches disable while master is off. Preference is device-scoped, not fridge-scoped.",
     "", "", ""),
    ("TC-PREF-002", "Notification settings", "Home badge switch zeroes badgeCount without marking items read", "High", "Functional",
     "Unread notifications exist.",
     "1. Disable home badges.\n2. Check Home and GET /api/notifications.",
     "badgeCount is 0 and the orange home number is hidden. unreadCount stays > 0. Inbox rows remain unread.",
     "", "", ""),
    ("TC-PREF-003", "Notification settings", "Quiet hours hide the home badge during the window", "High", "Functional",
     "Badges enabled; unread items exist.",
     "1. Enable quiet hours covering the current time.\n2. Check Home.\n3. Set the window so now is outside it.",
     "During quiet hours badgeCount is 0. Inbox history is unchanged. After the window the badge returns if unread items remain.",
     "", "", ""),
    ("TC-PREF-004", "Notification settings", "System delivery requests OS permission before enabling", "High", "Functional",
     "System notifications were previously off.",
     "1. Turn on system delivery.\n2. Deny OS permission, then allow it from Settings.",
     "Deny keeps systemDeliveryEnabled false and offers OS settings. Allow enables delivery. Expo Go may grant local reminders without a remote push token.",
     "", "", ""),
    ("TC-PREF-005", "Notification settings", "Type switches filter expiring, restock, shared, and system independently", "Medium", "Functional",
     "Each type can be produced in the inbox.",
     "1. Toggle shared off.\n2. Toggle restock off.\n3. Reload inbox.",
     "Hidden types are omitted from the list. Other types remain. Failed saves reload the server snapshot.",
     "", "", ""),

    # Local / push delivery
    ("TC-PUSH-001", "System reminders", "Local expiry reminders are scheduled 3 days before expiry", "High", "Functional",
     "OS permission granted; master, freshness, and system delivery on; item expires in more than 3 days.",
     "1. Save the item.\n2. Inspect scheduled local notifications if tooling allows, or wait/adjust device time in a lab.",
     "At most 32 upcoming dated batches get one local reminder each, timed 3 days before expiry and shifted out of quiet hours. Changing expiry or deleting the item cancels the old identifier.",
     "", "", ""),
    ("TC-PUSH-002", "System reminders", "Local inactivity reminder is pushed 7 days ahead on each active use", "Medium", "Functional",
     "System delivery and system category enabled.",
     "1. Use the app.\n2. Use it again the next day.",
     "A single 7-day inactivity reminder is rescheduled on activity so it only fires after genuine idle time.",
     "", "", ""),
    ("TC-PUSH-003", "System reminders", "Remote Expo push is sent to other shared members, not the actor", "High", "Functional",
     "Development/production build with Expo push token (not Expo Go if token cannot be issued). Shared + system delivery on. Not in quiet hours.",
     "1. Member A adds an item.\n2. Observe member B's OS notification.\n3. Observe A.",
     "B may receive an Expo push. A does not. A failed push does not roll back the inventory mutation. Delivery is audited as sent/failed/suppressed.",
     "", "", ""),
    ("TC-PUSH-004", "System reminders", "Quiet hours and missing token suppress remote delivery", "Medium", "Negative",
     "Shared fridge; B is in quiet hours or never registered a token.",
     "1. A updates inventory.",
     "B still gets the in-app shared notification. Remote push is suppressed. Inventory write succeeds.",
     "", "", ""),

    # Device recovery
    ("TC-DEV-001", "Device recovery", "Recovery code can be viewed in settings", "Medium", "Functional",
     "Device is bootstrapped.",
     "1. Open device recovery settings.",
     "A one-time recovery code (or the current code state) is available. The raw Device-Credential is not printed in logs or API responses.",
     "", "", ""),
    ("TC-DEV-002", "Device recovery", "Successful recovery transfers ownership and revokes the old device", "High", "Functional",
     "Old device has inventory; new device has the recovery code.",
     "1. Recover onto the new device.\n2. Use the old device afterwards.",
     "Ownership, membership, and notification reads move to the new device. Old credential is revoked. Recovery code rotates.",
     "", "", ""),
    ("TC-DEV-003", "Device recovery", "Wrong recovery code is limited and rejected", "Medium", "Negative",
     "New device recovery screen.",
     "1. Enter an incorrect code several times.",
     "Requests fail without transferring data. Repeated failures are limited.",
     "", "", ""),

    # i18n
    ("TC-I18N-001", "Language", "Interface language can switch between Chinese and English", "High", "Functional",
     "Me > Language is available.",
     "1. Switch to English.\n2. Inspect Home, Fridge, Add sheet, Inbox, and Profile.\n3. Switch back to Chinese.",
     "All user-facing strings update, including Use first / 先用, notification copy, and add-item labels. Preference persists after relaunch.",
     "", "", ""),
    ("TC-I18N-002", "Language", "Filters keep working after a language change", "Medium", "Functional",
     "Expiring filter is selected.",
     "1. Apply Expiring.\n2. Switch language.",
     "The same internal filter remains selected. Only labels change.",
     "", "", ""),

    # Sync / API
    ("TC-SYNC-001", "Sync", "Shared inventory change appears on the other member without reopening the app", "High", "Integration",
     "Two devices in one shared fridge; both on Fridge or Home.",
     "1. Device A adds or archives an item.\n2. Watch device B.",
     "B silently refreshes. Home expiring count and fridge list stay consistent. No business records are sent over Broadcast, only version invalidation.",
     "", "", ""),
    ("TC-SYNC-002", "Sync", "Returning to foreground refreshes home and fridge data", "Medium", "Integration",
     "App was backgrounded; inventory changed on another device or via API.",
     "1. Background the app.\n2. Change stock elsewhere.\n3. Foreground the app.",
     "Home count, fridge list, cart, and notifications catch up after the foreground probe.",
     "", "", ""),
    ("TC-SYNC-003", "Sync", "Health check drives the connection status copy", "Low", "Functional",
     "Express can be started and stopped.",
     "1. With API up, launch Home/placeholder connection text if visible.\n2. Stop Express and retry a data screen.",
     "Connected vs unavailable messaging is honest. Fridge/home counts fail softly rather than showing stale sample data as truth.",
     "", "", ""),
    ("TC-SYNC-004", "Sync", "API requires Device-ID and Device-Credential", "High", "Negative",
     "curl or HTTP client against Express.",
     "1. GET /api/inventory without headers.\n2. Repeat with only Device-ID.",
     "Unauthorised requests are rejected. Inventory is never returned to an unauthenticated caller.",
     "", "", ""),
    ("TC-SYNC-005", "Sync", "GET /api/health does not require device auth", "Low", "Functional",
     "Express is running.",
     "1. GET /api/health.",
     "Returns { status: ok, database: connected } when Supabase is reachable.",
     "", "", ""),
]


def style_header(ws, row=1):
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def autosize(ws, widths):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_cases(wb):
    ws = wb.active
    ws.title = "Test Cases"
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(CASES) + 1}"
    ws.row_dimensions[1].height = 24

    for r, case in enumerate(CASES, 2):
        for c, value in enumerate(case, 1):
            cell = ws.cell(r, c, value)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if r % 2 == 0:
                cell.fill = ALT_FILL
        priority = case[3]
        ws.cell(r, 4).alignment = CENTER
        ws.cell(r, 5).alignment = CENTER
        ws.cell(r, 1).alignment = CENTER
        ws.cell(r, 10).alignment = CENTER
        if priority == "High":
            ws.cell(r, 4).fill = HIGH_FILL
        elif priority == "Medium":
            ws.cell(r, 4).fill = MED_FILL
        else:
            ws.cell(r, 4).fill = LOW_FILL
        ws.row_dimensions[r].height = 78

    status_dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not executed,N/A"', allow_blank=True)
    status_dv.error = "Select a status from the list"
    status_dv.errorTitle = "Invalid status"
    status_dv.prompt = "Pass / Fail / Blocked / Not executed / N/A"
    status_dv.promptTitle = "Status"
    ws.add_data_validation(status_dv)
    status_dv.add(f"J2:J{len(CASES) + 1}")

    ws.conditional_formatting.add(
        f"J2:J{len(CASES) + 1}",
        FormulaRule(formula=['$J2="Pass"'], fill=PASS_FILL),
    )
    ws.conditional_formatting.add(
        f"J2:J{len(CASES) + 1}",
        FormulaRule(formula=['$J2="Fail"'], fill=FAIL_FILL),
    )
    ws.conditional_formatting.add(
        f"J2:J{len(CASES) + 1}",
        FormulaRule(formula=['$J2="Blocked"'], fill=BLOCK_FILL),
    )
    ws.conditional_formatting.add(
        f"J2:J{len(CASES) + 1}",
        FormulaRule(formula=['OR($J2="N/A",$J2="Not executed")'], fill=NA_FILL),
    )

    autosize(ws, [14, 18, 42, 12, 14, 36, 48, 48, 28, 14, 24])
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = "KitchMemo — Test Cases"
    return ws


def label_value(ws, row, label, value, value_col=3):
    ws.cell(row, 1, label).font = Font(name="Calibri", bold=True, size=11, color="173D31")
    ws.cell(row, 1).fill = LABEL_FILL
    ws.cell(row, 1).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell = ws.cell(row, value_col, value)
    cell.font = BODY_FONT
    cell.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=value_col, end_row=row, end_column=6)
    for col in range(1, 7):
        ws.cell(row, col).border = THIN
        ws.cell(row, col).alignment = WRAP


def build_report(wb, last_case_row):
    ws = wb.create_sheet("Test Report")
    ws.merge_cells("A1:F1")
    ws["A1"] = "KitchMemo Test Report"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"] = "Fill the Status column on Test Cases. Totals on this sheet update automatically."
    ws["A2"].font = Font(name="Calibri", italic=True, size=11, color="5B6B64")

    ws["A4"] = "1. Document information"
    ws["A4"].font = SECTION_FONT
    ws.merge_cells("A4:F4")

    info = [
        (5, "Product", "KitchMemo"),
        (6, "Version", "1.0.0 (Expo app + Express API + Supabase)"),
        (7, "Test round", "Functional test of currently implemented MVP features"),
        (8, "Test date", "2026-09-03"),
        (9, "Tester(s)", ""),
        (10, "Environment", "Physical Android/iOS device or Expo Go; Express on LAN; development Supabase project"),
        (11, "Build notes", "Chinese is the default UI language. English is switched from Me > Language. Recipes stay paused. Profile/Me is in the tab bar; Wins/achievements is still hidden."),
        (12, "In scope", "Onboarding, home 3D kitchen, freshness prompt, fridge inventory, add/manual/photo/AI preset, item detail, shopping checkout of checked items, in-app notifications, shared inventory events, device profile nickname, notification preferences, local expiry/inactivity reminders, Expo push to other members, sharing, device recovery, i18n, sync."),
        (13, "Out of scope", "Achievements API, category management API, explicit discard action, recipe discovery as a destination, Expo push receipt polling, user-uploaded avatars, persisted taste/diet preferences."),
    ]
    for row, label, value in info:
        label_value(ws, row, label, value)

    ws["A15"] = "2. Execution summary"
    ws["A15"].font = SECTION_FONT
    ws.merge_cells("A15:C15")

    summary_headers = ["Metric", "Count", "Notes"]
    for col, title in enumerate(summary_headers, 1):
        cell = ws.cell(16, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    rng = f"'Test Cases'!$J$2:$J${last_case_row}"
    rows = [
        (17, "Total cases", f"=COUNTA('Test Cases'!A2:A{last_case_row})", "All designed cases in this workbook"),
        (18, "Executed", f"=COUNTIF({rng},\"Pass\")+COUNTIF({rng},\"Fail\")", "Pass + Fail"),
        (19, "Pass", f"=COUNTIF({rng},\"Pass\")", "Behaviour matches expected result"),
        (20, "Fail", f"=COUNTIF({rng},\"Fail\")", "Defect confirmed"),
        (21, "Blocked", f"=COUNTIF({rng},\"Blocked\")", "Cannot run due to environment or dependency"),
        (22, "Not executed", f"=COUNTIF({rng},\"Not executed\")+COUNTBLANK({rng})", "Blank status counts as not executed"),
        (23, "N/A", f"=COUNTIF({rng},\"N/A\")", "Not applicable in this environment"),
        (24, "Pass rate", f'=IF(B18=0,"N/A",B19/B18)', "Pass / Executed"),
    ]
    for row, metric, formula, note in rows:
        ws.cell(row, 1, metric).font = BODY_FONT
        ws.cell(row, 1).border = THIN
        ws.cell(row, 1).fill = LABEL_FILL
        cell = ws.cell(row, 2, formula)
        cell.font = Font(name="Calibri", bold=True, size=11)
        cell.border = THIN
        cell.alignment = CENTER
        ws.cell(row, 3, note).font = BODY_FONT
        ws.cell(row, 3).border = THIN
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
    ws["B24"].number_format = "0.0%"
    ws["B19"].fill = PASS_FILL
    ws["B20"].fill = FAIL_FILL
    ws["B21"].fill = BLOCK_FILL

    ws["A26"] = "3. Coverage by module"
    ws["A26"].font = SECTION_FONT
    ws.merge_cells("A26:F26")

    module_headers = ["Module", "Total", "Pass", "Fail", "Blocked", "Not executed"]
    for col, title in enumerate(module_headers, 1):
        cell = ws.cell(27, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    modules = sorted({case[1] for case in CASES})
    for i, module in enumerate(modules):
        row = 28 + i
        ws.cell(row, 1, module).border = THIN
        ws.cell(row, 1).font = BODY_FONT
        ws.cell(row, 2, f'=COUNTIF(\'Test Cases\'!B:B,A{row})').border = THIN
        ws.cell(row, 3, f'=COUNTIFS(\'Test Cases\'!B:B,A{row},\'Test Cases\'!J:J,"Pass")').border = THIN
        ws.cell(row, 4, f'=COUNTIFS(\'Test Cases\'!B:B,A{row},\'Test Cases\'!J:J,"Fail")').border = THIN
        ws.cell(row, 5, f'=COUNTIFS(\'Test Cases\'!B:B,A{row},\'Test Cases\'!J:J,"Blocked")').border = THIN
        ws.cell(row, 6, f'=COUNTIFS(\'Test Cases\'!B:B,A{row},\'Test Cases\'!J:J,"Not executed")+COUNTIFS(\'Test Cases\'!B:B,A{row},\'Test Cases\'!J:J,"")').border = THIN
        for col in range(2, 7):
            ws.cell(row, col).alignment = CENTER
            ws.cell(row, col).font = BODY_FONT
        ws.cell(row, 3).fill = PASS_FILL
        ws.cell(row, 4).fill = FAIL_FILL

    last_module_row = 27 + len(modules)
    pie_start = last_module_row + 2
    ws.cell(pie_start, 1, "4. Result distribution (chart data)").font = SECTION_FONT
    ws.merge_cells(start_row=pie_start, start_column=1, end_row=pie_start, end_column=2)
    ws.cell(pie_start + 1, 1, "Result")
    ws.cell(pie_start + 1, 2, "Count")
    for col in (1, 2):
        ws.cell(pie_start + 1, col).fill = HEADER_FILL
        ws.cell(pie_start + 1, col).font = HEADER_FONT
        ws.cell(pie_start + 1, col).alignment = CENTER
        ws.cell(pie_start + 1, col).border = THIN
    chart_rows = [
        (pie_start + 2, "Pass", "=B19"),
        (pie_start + 3, "Fail", "=B20"),
        (pie_start + 4, "Blocked", "=B21"),
        (pie_start + 5, "Not executed", "=B22"),
        (pie_start + 6, "N/A", "=B23"),
    ]
    for row, name, formula in chart_rows:
        ws.cell(row, 1, name).border = THIN
        ws.cell(row, 2, formula).border = THIN
        ws.cell(row, 2).alignment = CENTER

    chart = PieChart()
    chart.title = "Test result distribution"
    labels = Reference(ws, min_col=1, min_row=pie_start + 2, max_row=pie_start + 6)
    data = Reference(ws, min_col=2, min_row=pie_start + 1, max_row=pie_start + 6)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = False
    chart.dataLabels.showCatName = True
    chart.width = 14
    chart.height = 8
    ws.add_chart(chart, "D16")

    findings_row = pie_start + 8
    ws.cell(findings_row, 1, "5. Findings and sign-off").font = SECTION_FONT
    ws.merge_cells(start_row=findings_row, start_column=1, end_row=findings_row, end_column=6)
    blocks = [
        (findings_row + 1, "Critical defects", ""),
        (findings_row + 3, "Major defects", ""),
        (findings_row + 5, "Minor defects / UI polish", ""),
        (findings_row + 7, "Overall conclusion", "Ready for course demo / Needs retest / Not ready  (choose one after execution)"),
        (findings_row + 9, "Tester sign-off", "Name:                    Date:"),
        (findings_row + 11, "Reviewer sign-off", "Name:                    Date:"),
    ]
    for row, label, value in blocks:
        label_value(ws, row, label, value)
        ws.row_dimensions[row].height = 36
        ws.row_dimensions[row + 1].height = 8

    ws.cell(findings_row + 13, 1, "How to use").font = SECTION_FONT
    ws.merge_cells(start_row=findings_row + 13, start_column=1, end_row=findings_row + 13, end_column=6)
    ws.merge_cells(start_row=findings_row + 14, start_column=1, end_row=findings_row + 16, end_column=6)
    ws.cell(
        findings_row + 14,
        1,
        "1. Execute each case on Test Cases and choose Status from the dropdown.\n"
        "2. Record Actual Result and Notes when a case fails or is blocked.\n"
        "3. Return here: summary counts, module coverage, and the pie chart refresh from column Status.\n"
        "4. Expiring window used by Home, Fridge Expiring chip, notification sync, and local OS expiry reminders is 3 days, not yet expired.\n"
        "5. Home orange mail number is badgeCount (respects badge switch and quiet hours). Inbox unreadCount is the true unread total for allowed types.\n"
        "6. Do not treat GET /api/health as an inventory save. Inventory writes go to /api/inventory/batches.\n"
        "7. Expo Go may allow local reminders but often cannot register a remote Expo push token; mark TC-PUSH-003 N/A in that environment.",
    ).alignment = WRAP
    ws.cell(findings_row + 14, 1).font = BODY_FONT
    ws.row_dimensions[findings_row + 14].height = 120

    autosize(ws, [22, 14, 18, 14, 14, 22])
    ws.row_dimensions[12].height = 48
    ws.row_dimensions[13].height = 48
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddHeader.left.text = "KitchMemo — Test Report"
    ws.print_area = f"A1:F{findings_row + 16}"
    return ws


def main():
    wb = Workbook()
    build_cases(wb)
    build_report(wb, last_case_row=len(CASES) + 1)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} with {len(CASES)} test cases")


if __name__ == "__main__":
    main()
